import docx, os, glob

files = sorted(glob.glob('E:/项目/归星纪/未定义文明/章节正文/*.docx'))
for f in files[-2:]:
    doc = docx.Document(f)
    print(f'=== {os.path.basename(f)} ({len(doc.paragraphs)}段) ===')
    for p in doc.paragraphs[:15]:
        if p.text.strip():
            print(p.text[:400])
    print()
