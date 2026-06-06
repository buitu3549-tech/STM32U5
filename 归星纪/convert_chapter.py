import re, os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def md_to_docx(md_path, docx_path, chapter_title, chapter_subtitle="归星纪 · 第一部 · 未定义文明"):
    with open(md_path, 'r', encoding='utf-8') as f:
        text = f.read()

    # Remove the markdown H1 title line and trailing markers
    body = re.sub(r'^#\s*第\d+章\s*.*?\n', '', text, count=1)
    body = re.sub(r'\n---\n\n\*\(第\d+章 完\)\*', '', body)
    body = re.sub(r'\n---\n\n\*\(.*?完\)\*', '', body)
    body = body.strip()

    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(3.18)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(chapter_title)
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.name = '宋体'
    run._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '宋体')

    # Subtitle
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(chapter_subtitle)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '宋体')

    doc.add_paragraph()

    # Body - split on blank lines
    paragraphs = body.split('\n\n')
    for p_text in paragraphs:
        p_text = p_text.strip()
        if not p_text:
            continue
        for sub_line in p_text.split('\n'):
            sub_line = sub_line.strip()
            if not sub_line:
                continue
            para = doc.add_paragraph()
            para.paragraph_format.first_line_indent = Pt(24)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.5
            run = para.add_run(sub_line)
            run.font.size = Pt(12)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '宋体')

    doc.add_paragraph()
    end_para = doc.add_paragraph()
    end_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = end_para.add_run(f'（{chapter_title} 完）')
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '宋体')

    doc.save(docx_path)
    print(f'Done: {docx_path}')

def convert_chapter(chapter_num, chapter_slug, docx_title=None):
    base = 'E:/项目/归星纪/未定义文明/章节正文'
    num_str = f'{chapter_num:03d}'
    md_path = f'{base}/第{num_str}章-{chapter_slug}.md'
    docx_path = f'{base}/第{num_str}章-{chapter_slug}.docx'
    if docx_title is None:
        docx_title = f'第{num_str}章　{chapter_slug}'
    md_to_docx(md_path, docx_path, docx_title)
    return docx_path

if __name__ == '__main__':
    if len(sys.argv) >= 3:
        chapter_num = int(sys.argv[1])
        chapter_slug = sys.argv[2]
        docx_title = sys.argv[3] if len(sys.argv) > 3 else None
        convert_chapter(chapter_num, chapter_slug, docx_title)
    else:
        # Default: convert chapter 1
        convert_chapter(1, '放学', '第001章　放学')
