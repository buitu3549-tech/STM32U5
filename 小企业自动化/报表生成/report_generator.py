"""
报表生成模板：Excel 汇总表 + Word 发货单 + PDF 报表
依赖：pandas, openpyxl, python-docx, weasyprint (可选)
"""

from pathlib import Path
from datetime import datetime, timedelta
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# Excel 报表：自动生成日报/周报/月报
# ============================================================

def generate_sales_report(df: pd.DataFrame, period: str = "daily", output_dir: str = "E:/报表/"):
    """
    df 需包含列：日期, 订单号, 商品, 数量, 单价, 金额, 客户, 地区
    自动按 period 汇总并生成 Excel
    """
    df["日期"] = pd.to_datetime(df["日期"])
    today = datetime.today()

    if period == "daily":
        mask = df["日期"].dt.date == today.date()
        filename = f"日报_{today.strftime('%Y%m%d')}.xlsx"
    elif period == "weekly":
        monday = today - timedelta(days=today.weekday())
        mask = df["日期"] >= monday
        filename = f"周报_{monday.strftime('%Y%m%d')}.xlsx"
    elif period == "monthly":
        mask = df["日期"].dt.month == today.month
        filename = f"月报_{today.strftime('%Y%m')}.xlsx"

    period_df = df[mask].copy()
    output = Path(output_dir) / filename
    output.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Sheet 1: 明细
        period_df.to_excel(writer, sheet_name="订单明细", index=False)

        # Sheet 2: 按商品汇总
        product_summary = period_df.groupby("商品").agg(
            销量=("数量", "sum"),
            销售额=("金额", "sum"),
        ).reset_index()
        product_summary.to_excel(writer, sheet_name="商品汇总", index=False)

        # Sheet 3: 按地区汇总
        if "地区" in df.columns:
            region_summary = period_df.groupby("地区").agg(
                订单数=("订单号", "nunique"),
                销售额=("金额", "sum"),
            ).reset_index()
            region_summary.to_excel(writer, sheet_name="地区汇总", index=False)

    print(f"报表已生成: {output}")
    return output


# ============================================================
# Word 发货单：从 CSV 批量生成快递单 / 发货单
# ============================================================

def generate_delivery_notes(orders_file: str, template_word: str = None, output_dir: str = "E:/发货单/"):
    """
    读取订单 CSV，为每条订单生成一个 Word 发货单
    CSV 列：订单号, 收货人, 电话, 地址, 商品清单, 备注
    """
    df = pd.read_csv(orders_file)
    output_dir = Path(output_dir) / datetime.today().strftime("%Y%m%d")
    output_dir.mkdir(parents=True, exist_ok=True)

    generated = []
    for _, row in df.iterrows():
        doc = Document()

        # 页面设置
        section = doc.sections[0]
        section.page_width = Cm(10)
        section.page_height = Cm(14)

        # 标题
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("发 货 单")
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.name = "宋体"

        # 订单信息
        info = doc.add_paragraph()
        info.paragraph_format.line_spacing = 1.5
        for label, key in [("订单号", "订单号"), ("收货人", "收货人"),
                           ("电话", "电话"), ("地址", "地址"), ("备注", "备注")]:
            if key in row and pd.notna(row[key]):
                line = info.add_run(f"\n{label}：{row[key]}")
                line.font.size = Pt(12)
                line.font.name = "宋体"

        # 保存
        order_id = str(row.get("订单号", _))
        filepath = output_dir / f"发货单_{order_id}.docx"
        doc.save(str(filepath))
        generated.append(filepath)

    print(f"已生成 {len(generated)} 份发货单 → {output_dir}")
    return generated


# ============================================================
# PDF 报表：HTML 转 PDF（精美排版）
# ============================================================

def generate_pdf_report(title: str, df: pd.DataFrame, output_path: str):
    """将 DataFrame 渲染为带 CSS 的精美 PDF 报表"""
    try:
        from weasyprint import HTML
    except ImportError:
        print("需要安装 weasyprint: pip install weasyprint")
        return None

    html = f"""
    <html><head><meta charset="utf-8">
    <style>
        body {{ font-family: 'Microsoft YaHei', sans-serif; padding: 40px; }}
        h1 {{ text-align: center; color: #333; }}
        table {{ width: 100%; border-collapse: collapse; margin-top: 20px; }}
        th {{ background: #4a90d9; color: white; padding: 10px; }}
        td {{ padding: 8px; border-bottom: 1px solid #ddd; text-align: center; }}
        tr:nth-child(even) {{ background: #f5f5f5; }}
        .summary {{ background: #e8f4f8; padding: 15px; margin: 20px 0; border-radius: 8px; }}
    </style></head><body>
        <h1>{title}</h1>
        <p style="text-align:center;color:#999;">生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}</p>
        <div class="summary">总订单数：{len(df)} | 总金额：¥{df['金额'].sum():,.2f}</div>
        {df.to_html(index=False, border=0)}
    </body></html>
    """

    HTML(string=html).write_pdf(output_path)
    print(f"PDF 已生成: {output_path}")
    return output_path


# ============================================================
# 快捷入口：从命令行一键调用
# ============================================================

if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("用法:")
        print("  python report_generator.py daily   数据.csv   → 生成日报")
        print("  python report_generator.py delivery 订单.csv → 生成发货单")
        sys.exit(0)

    mode = sys.argv[1]

    if mode == "daily":
        df = pd.read_csv(sys.argv[2])
        generate_sales_report(df, "daily")
    elif mode == "delivery":
        generate_delivery_notes(sys.argv[2])
